"""Extract labels/styles only; never export financial amounts or memo columns."""
import json, sys, uuid, re, hashlib
from pathlib import Path
import openpyxl
path = Path(sys.argv[1]) if len(sys.argv)>1 else Path('/workspace/attachments/f7efe7a1-62e5-4b32-bc49-3b0f9cf2cc22/GALAPAGOS 売上 経費.xlsx')
w=openpyxl.load_workbook(path, data_only=False)
ns=uuid.UUID('be888791-cd03-4dc4-98ac-89079f87d506')
def uid(s):return str(uuid.uuid5(ns,s))
def color(c):
 x=c.fill.fgColor
 return {'type':x.type,'value':str(x.rgb if x.type=='rgb' else x.indexed if x.type=='indexed' else x.theme),'tint':x.tint,'pattern':c.fill.patternType}
categories=[]; report=[]
for year in [2025,2024,2023]:
 s=w[f'{year}年度　経費']; col=1 if year==2025 else 3
 labels=[(r,str(s.cell(r,col).value).strip()) for r in range(2 if year==2025 else 4,s.max_row+1) if s.cell(r,col).value is not None]
 totals=[(r,n) for r,n in labels if '合計' in n]
 groups={}; assignment={}
 for r,n in totals:
  # Contiguous named rows above a subtotal, including rows accidentally omitted by its formula.
  start=r-1
  while start>1 and (s.cell(start,col).value is not None or s.cell(start-1,col).value is not None):
   if s.cell(start,col).value and '合計' in str(s.cell(start,col).value):break
   start-=1
  start+=1
  gid=uid(f'{year}:group:{r}'); groups[gid]=(n.replace('合計','').strip(),r)
  for rr,nn in labels:
   if start<=rr<r and '合計' not in nn:assignment[rr]=gid
 for r,n in labels:
  if '合計' in n: continue
  gid=assignment.get(r)
  if not gid:
   gid=uid(f'{year}:group:standalone:{r}');groups[gid]=(n,r)
  categories.append({'id':uid(f'{year}:{col}:{r}'),'parent_id':gid,'name':n,'year':year,'source_sheet':s.title,'source_cell':s.cell(r,col).coordinate,'order':r,'selectable':True,'color':color(s.cell(r,col))})
 for gid,(n,r) in groups.items():categories.append({'id':gid,'parent_id':None,'name':n,'year':year,'source_sheet':s.title,'source_cell':s.cell(r,col).coordinate,'order':r,'selectable':False,'color':color(s.cell(r,col))})
 report.append({'year':year,'labels':len(labels),'categories':len(labels)-len(totals),'subtotals':len(totals),'excluded_total_cells':[s.cell(r,col).coordinate for r,n in totals]})
Path('data/categories.json').write_text(json.dumps(categories,ensure_ascii=False,indent=2)+'\n')
Path('data/excel-provenance.json').write_text(json.dumps({'file_name':path.name,'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'sheets':report,'notes':['分類ラベルと書式のみ移行。金額・売上・メモは取り込まない。','年度とセル座標で固有IDを生成。同名を統合しない。','小計直前の連続項目をグループ化。元の数式で抜けた項目も保持。','2023年など小計のない項目は独立グループ。用途への再編はユーザー承認が必要。']},ensure_ascii=False,indent=2)+'\n')
print(json.dumps(report,ensure_ascii=False))
